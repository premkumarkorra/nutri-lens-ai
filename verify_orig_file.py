import docx

orig_path = r"C:\Users\premc\.gemini\antigravity\brain\30231225-664b-4840-ba1d-8ecb4859952f\.user_uploaded\media_1792817472657_7d1aafa1.docx"
doc = docx.Document(orig_path)
print(f"File exists and is readable: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} sections")
print("Page 1 Title paragraph:", doc.paragraphs[0].text[:80])
print("P139 (Declaration):", doc.paragraphs[139].text[:80])
print("P149 (Last front matter):", doc.paragraphs[149].text[:80])
